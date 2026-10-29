"""Banksalad '엑셀 내보내기' parser.

The export is one workbook with two sheets:
  - '가계부 내역': flat ledger — 날짜, 시간, 타입, 대분류, 소분류, 내용, 금액, 화폐, 결제수단, 메모
  - a summary sheet whose '3.재무현황' block lists asset/debt balances at export time
"""

import hashlib
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, time
from io import BytesIO

import openpyxl

LEDGER_SHEET = "가계부 내역"
HEADER = ("날짜", "시간", "타입", "대분류", "소분류", "내용", "금액", "화폐", "결제수단", "메모")
_PERIOD = re.compile(r"(\d{4}-\d{2}-\d{2})~(\d{4}-\d{2}-\d{2})")


class ExportFormatError(ValueError):
    pass


@dataclass(frozen=True)
class Row:
    occurred_at: datetime
    src_type: str
    src_cat: str
    src_sub: str
    merchant: str
    amount: int
    currency: str
    account: str
    memo: str | None


@dataclass
class Balance:
    side: str  # asset | debt
    group: str
    name: str
    amount: int


@dataclass
class Investment:
    product_type: str
    company: str
    name: str
    cost: int
    value: int


@dataclass
class Export:
    rows: list[Row]
    skipped: int  # zero-amount rows (Banksalad logs 0원 interest entries)
    period_start: date | None
    period_end: date | None
    balances: list[Balance] = field(default_factory=list)
    investments: list[Investment] = field(default_factory=list)


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _when(d, t) -> datetime:
    if isinstance(d, datetime):
        day = d.date()
    elif isinstance(d, date):
        day = d
    else:
        day = datetime.strptime(_text(d)[:10], "%Y-%m-%d").date()
    if isinstance(t, datetime):
        t = t.time()
    if not isinstance(t, time):
        s = _text(t)
        t = datetime.strptime(s, "%H:%M:%S").time() if s else time(0, 0)
    return datetime.combine(day, t)


def _amount(v) -> int:
    if isinstance(v, (int, float)):
        return round(v)
    return round(float(_text(v).replace(",", "") or 0))


def parse_ledger(ws) -> tuple[list[Row], int]:
    it = ws.iter_rows(values_only=True)
    header = tuple(_text(v) for v in next(it, ()))[: len(HEADER)]
    if header != HEADER:
        raise ExportFormatError(f"'{LEDGER_SHEET}' 시트의 열 구성이 예상과 달라요: {header}")
    rows: list[Row] = []
    skipped = 0
    for r in it:
        if not r or r[0] is None:
            continue
        amount = _amount(r[6])
        if amount == 0:
            skipped += 1
            continue
        rows.append(
            Row(
                occurred_at=_when(r[0], r[1]),
                src_type=_text(r[2]),
                src_cat=_text(r[3]),
                src_sub=_text(r[4]),
                merchant=_text(r[5]),
                amount=amount,
                currency=_text(r[7]) or "KRW",
                account=_text(r[8]),
                memo=_text(r[9]) or None,
            )
        )
    return rows, skipped


def parse_balances(ws) -> list[Balance]:
    """Read the '3.재무현황' block: assets on the left (항목/상품명/·/금액), debts on the right."""
    out: list[Balance] = []
    in_block = False
    started = False
    asset_group = debt_group = ""
    for r in ws.iter_rows(values_only=True):
        cells = [_text(v) for v in r] + [""] * 10
        raw = list(r) + [None] * 10
        first = cells[1]
        if first.startswith("3."):
            in_block = True
            continue
        if not in_block:
            continue
        if first[:2] in ("4.", "5.", "6.") or first == "총자산":
            break
        if first == "항목" and cells[2] == "상품명":
            started = True
            continue
        if not started:
            continue
        if first:
            asset_group = first
        if cells[2] and isinstance(raw[4], (int, float)):
            out.append(Balance("asset", asset_group, cells[2], round(raw[4])))
        if cells[5]:
            debt_group = cells[5]
        if cells[6] and isinstance(raw[8], (int, float)):
            out.append(Balance("debt", debt_group, cells[6], round(raw[8])))
    return out


def parse_investments(ws) -> list[Investment]:
    """Read the '5.투자현황' block: 투자상품종류 | 금융사 | 상품명 | · | 투자원금 | 평가금액 | 수익률 …"""
    out: list[Investment] = []
    in_block = started = False
    for r in ws.iter_rows(values_only=True):
        cells = [_text(v) for v in r] + [""] * 10
        raw = list(r) + [None] * 10
        first = cells[1]
        if first.startswith("5."):
            in_block = True
            continue
        if not in_block:
            continue
        if first[:2] in ("6.", "7."):
            break
        if first == "투자상품종류":
            started = True
            continue
        if not started or first == "총계" or not cells[3]:
            continue
        if cells[3].startswith("총 ") or not isinstance(raw[6], (int, float)):
            continue
        cost = raw[5] if isinstance(raw[5], (int, float)) else 0
        out.append(Investment(first, cells[2], cells[3], round(cost), round(raw[6])))
    return out


def parse_export(data: bytes, filename: str = "") -> Export:
    try:
        wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ExportFormatError("엑셀 파일을 열 수 없어요. 암호가 걸려 있다면 풀고 올려 주세요.") from e
    if LEDGER_SHEET not in wb.sheetnames:
        raise ExportFormatError(f"'{LEDGER_SHEET}' 시트가 없어요. 뱅크샐러드 내보내기 파일이 맞나요?")
    rows, skipped = parse_ledger(wb[LEDGER_SHEET])
    balances: list[Balance] = []
    investments: list[Investment] = []
    for name in wb.sheetnames:
        if name != LEDGER_SHEET:
            balances = parse_balances(wb[name])
            investments = parse_investments(wb[name])
            break

    m = _PERIOD.search(filename)
    if m:
        start = date.fromisoformat(m.group(1))
        end = date.fromisoformat(m.group(2))
    elif rows:
        start = min(r.occurred_at for r in rows).date()
        end = max(r.occurred_at for r in rows).date()
    else:
        start = end = None
    return Export(
        rows=rows, skipped=skipped, period_start=start, period_end=end, balances=balances,
        investments=investments,
    )


def base_key(row: Row) -> str:
    """Identity of a ledger row across exports.

    Deliberately excludes Banksalad's 타입/분류, which it re-labels over time
    (e.g. card bill payments moved from 이체 to 지출 in 2026).
    """
    raw = "|".join(
        [row.account, row.occurred_at.isoformat(), str(row.amount), row.currency, row.merchant]
    )
    return hashlib.sha1(raw.encode()).hexdigest()


def row_keys(rows: list[Row]) -> list[tuple[str, int]]:
    """(base_key, occurrence) per row; the nth identical row in an export gets occurrence n."""
    seen: Counter[str] = Counter()
    keys = []
    for row in rows:
        k = base_key(row)
        seen[k] += 1
        keys.append((k, seen[k]))
    return keys
